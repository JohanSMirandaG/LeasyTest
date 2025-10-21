import io
from django.views.generic import ListView
from django.db.models import Q
from cars.models import Car
from .models import Contract
from django.contrib.auth.mixins import LoginRequiredMixin
from datetime import datetime
from django.contrib import messages
from django.db import transaction
from django.shortcuts import render, redirect
from django.views import View
import pandas as pd
from .forms import UploadFileForm
from clients.models import Client
import pytz
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

class DashboardView(LoginRequiredMixin, ListView):
    template_name = "contracts/dashboard.html"
    context_object_name = "contracts"
    paginate_by = 20
    login_url = "/accounts/login/"
    redirect_field_name = "next"

    def get_queryset(self):
        queryset = (
            Contract.objects
            .select_related("client", "car")
            .filter(is_active=True)
            .order_by("-created_at")
        )

        query = self.request.GET.get("q")
        if query:
            queryset = queryset.filter(
                Q(client__first_name__icontains=query)
                | Q(client__last_name__icontains=query)
                | Q(client__document_number__icontains=query)
                | Q(car__plate__icontains=query)
                | Q(car__brand__icontains=query)
                | Q(car__model__icontains=query)
            )
        return queryset

class ContractUploadView(LoginRequiredMixin, View):
    template_name = "contracts/upload.html"
    required_columns = [
        "Nombres",
        "Apellidos",
        "Número de documento",
        "Inicio de contrato",
        "Cuota semanal",
        "Marca del auto",
        "Modelo del auto",
        "Placa del auto",
    ]

    def get(self, request):
        form = UploadFileForm()
        return render(request, self.template_name, {"form": form})

    def post(self, request):
        form = UploadFileForm(request.POST, request.FILES)
        if not form.is_valid():
            messages.error(request, "Debe seleccionar un archivo válido.")
            return render(request, self.template_name, {"form": form})

        file = request.FILES["file"]
        filename = file.name.lower()

        if not (filename.endswith(".xlsx") or filename.endswith(".csv")):
            messages.error(request, "Formato no permitido. Solo se aceptan archivos .xlsx o .csv.")
            return render(request, self.template_name, {"form": form})

        # --- Leer el archivo ---
        try:
            content = file.read()
            buffer = io.BytesIO(content)

            try:
                df = pd.read_excel(buffer)
            except Exception:
                buffer.seek(0)
                sep = "," if content[:4096].count(b",") > content[:4096].count(b";") else ";"
                df = pd.read_csv(io.StringIO(content.decode("utf-8")), sep=sep)

        except Exception as e:
            messages.error(request, f"No se pudo leer el archivo: {e}")
            return render(request, self.template_name, {"form": form})

        # Normaliza columnas
        df.columns = [str(c).strip() for c in df.columns]
        missing = [c for c in self.required_columns if c not in df.columns]
        if missing:
            messages.error(request, f"Faltan columnas requeridas: {', '.join(missing)}")
            return render(request, self.template_name, {"form": form})

        rows = df.to_dict(orient="records")
        row_errors = []

        # --- Validar campos requeridos ---
        for idx, row in enumerate(rows, start=2):
            missing_fields = [
                c for c in self.required_columns
                if not row.get(c) or str(row.get(c)).strip() == "" or pd.isna(row.get(c))
            ]
            if missing_fields:
                row_errors.append({"row": idx, "missing": missing_fields})

        if row_errors:
            context = {
                "form": form,
                "validation_errors": row_errors,
                "num_rows": len(rows),
            }
            messages.error(request,
                           "Some rows contain missing or invalid required fields. Please correct and re-upload.")
            return render(request, self.template_name, context)

        # --- Procesar datos ---
        created_clients = created_cars = created_contracts = 0
        updated_contracts = 0
        bogota_tz = pytz.timezone("America/Bogota")

        try:
            with transaction.atomic():
                for row in rows:
                    first_name = str(row["Nombres"]).strip()
                    last_name = str(row["Apellidos"]).strip()
                    document = str(row["Número de documento"]).strip()
                    start_raw = row["Inicio de contrato"]
                    weekly_raw = row["Cuota semanal"]
                    brand = str(row["Marca del auto"]).strip()
                    model = str(row["Modelo del auto"]).strip()
                    plate = str(row["Placa del auto"]).strip()

                    # Fecha de inicio
                    try:
                        if isinstance(start_raw, (pd.Timestamp, datetime)):
                            start_date = start_raw.date()
                        else:
                            start_date = pd.to_datetime(start_raw).date()
                    except Exception:
                        start_date = datetime.now().date()

                    weekly_amount = float(weekly_raw)

                    # Cliente
                    client_obj, client_created = Client.objects.get_or_create(
                        document_number=document,
                        defaults={
                            "first_name": first_name,
                            "last_name": last_name,
                            "registration_date": datetime.now(bogota_tz).date(),
                        },
                    )
                    if client_created:
                        created_clients += 1
                    else:
                        updated = False
                        if first_name and client_obj.first_name != first_name:
                            client_obj.first_name = first_name
                            updated = True
                        if last_name and client_obj.last_name != last_name:
                            client_obj.last_name = last_name
                            updated = True
                        if updated:
                            client_obj.save()

                    # --- Auto ---
                    car_obj, car_created = Car.objects.get_or_create(
                        plate=plate,
                        defaults={
                            "brand": brand,
                            "model": model,
                            "fabrication_date": datetime.now(bogota_tz).date(),
                        },
                    )
                    if car_created:
                        created_cars += 1

                    # --- Contrato ---
                    active_client_contract = Contract.objects.filter(client=client_obj, is_active=True).first()
                    active_car_contract = Contract.objects.filter(car=car_obj, is_active=True).first()

                    # Si el cliente ya tiene contrato activo o el auto está en uso → inactivar
                    if active_client_contract:
                        active_client_contract.is_active = False
                        active_client_contract.save()
                        updated_contracts += 1

                    if active_car_contract:
                        active_car_contract.is_active = False
                        active_car_contract.save()
                        updated_contracts += 1

                    # Crear nuevo contrato
                    Contract.objects.create(
                        client=client_obj,
                        car=car_obj,
                        weekly_amount=weekly_amount,
                        weeks=52,
                        start_date=start_date,
                        is_active=True,
                    )
                    created_contracts += 1

        except Exception as e:
            messages.error(request, f"Error al guardar los datos: {e}")
            return render(request, self.template_name, {"form": form})

        messages.success(request,
            f"Importación completada: {len(rows)} filas procesadas — "
            f"{created_clients} nuevos clientes, {created_cars} nuevos autos, "
            f"{created_contracts} contratos creados, {updated_contracts} contratos actualizados/inactivados."
        )
        return redirect("dashboard")

class ContractReportView(LoginRequiredMixin, View):
    def post(self, request):
        selected_fields = request.POST.getlist("fields")
        status_filter = request.POST.get("status", "active")

        if not selected_fields:
            messages.error(request, "Debe seleccionar al menos una columna para generar el reporte.")
            return redirect("dashboard")

        contracts = Contract.objects.select_related("client", "car").order_by("id")

        if status_filter == "active":
            contracts = contracts.filter(is_active=True)
        elif status_filter == "inactive":
            contracts = contracts.filter(is_active=False)

        field_mapping = {
            "id": ("ID", lambda c: c.id),
            "client_name": ("Cliente", lambda c: f"{c.client.first_name} {c.client.last_name}"),
            "document": ("Documento", lambda  c: c.client.document_number),
            "car": ("Auto", lambda c: f"{c.car.plate} - {c.car.brand} {c.car.model}"),
            "weekly_amount": ("Valor semanal", lambda c: float(c.weekly_amount)),
            "weeks": ("Semanas", lambda c: c.weeks),
            "start_date": ("Fecha inicio", lambda c: c.start_date.strftime("%Y-%m-%d")),
            "created_at": ("Creado", lambda c: c.created_at.strftime("%Y-%m-%d %H:%M")),
            "updated_at": ("Actualizado", lambda c: c.updated_at.strftime("%Y-%m-%d %H:%M")),
        }

        wb = Workbook()
        ws = wb.active
        ws.title = "Contratos"

        headers = [field_mapping[f][0] for f in selected_fields]
        ws.append(headers)

        for contract in contracts:
            row = [field_mapping[f][1](contract) for f in selected_fields]
            ws.append(row)

        for i, _ in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width = 20

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        report_name = f"contracts_report_{status_filter}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{report_name}"'
        wb.save(response)
        return response